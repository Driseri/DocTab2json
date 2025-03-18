from docx import Document
import json
from typing import Dict, List, Optional
import sys

def is_bold(paragraph) -> bool:
    """ЖИРНЫЙ КТО-то?"""
    return any(run.bold for run in paragraph.runs if run.bold is not None)


def docx_table_to_json(docx_path: str, output_path: Optional[str] = None) -> Dict[str, List[str]]:
    """
    :param docx_path: Путь к файлу .docx
    :param output_path: Если указан, сохраняет результат в JSON-файл
    :return: Словарь, где жирные ячейки - ключи, обычные ячейки - список значений
    """
    try:
        doc = Document(docx_path)
    except Exception as e:
        raise ValueError(f"Ошибка при открытии документа: {e}")

    data = {}  # Итоговый JSON-словарь
    current_key = None  # Текущий атрибут
    values = []  # Список значений для текущего атрибута

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) > 1:
                cell = row.cells[1]
                if any(is_bold(para) for para in cell.paragraphs):
                    if current_key and values:
                        data[current_key] = values
                    current_key = cell.text.strip()
                    values = []
                else:
                    if cell.text.strip():
                        values.append(cell.text.strip())

    if current_key and values:
        data[current_key] = values

    if output_path:
        try:
            with open(output_path, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=4, ensure_ascii=False)
            print(f"Данные успешно сохранены в {output_path}")
        except Exception as e:
            raise ValueError(f"Ошибка при сохранении JSON: {e}")

    return data


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Использование: python table2json.py <путь_к_файлу.docx> <опционально: путь_к_JSON>")
    else:
        docx_file = sys.argv[1]
        json_output = sys.argv[2] if len(sys.argv) > 2 else None
        docx_table_to_json(docx_file, json_output)
